# -*- coding: utf-8 -*-
"""把三篇学习心得(.md) 清理脚手架并转成 Word。
清理：去斜体提示语/本文结构/字数标记/空日期/正文加粗；填姓名与研读日期。
运行：py build_reflections_docx.py
"""
import re
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 数模/
OUT = os.path.join(BASE, '学习心得_word')
os.makedirs(OUT, exist_ok=True)

# ---------------- 数学公式渲染（matplotlib mathtext → 内嵌图）----------------
import struct
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams['mathtext.fontset'] = 'cm'   # 经典 LaTeX 风格
_FX = os.path.join(OUT, '_fx')
os.makedirs(_FX, exist_ok=True)
_mcache = {}

# 纯文本公式 → LaTeX（按原文精确串匹配，替换为排版图；未命中则保持原文）
FORMULAS = {
    'R = |AB| / (2·sin α)': r'R=\frac{|AB|}{2\sin\alpha}',
    '|AB| / sin(∠APB) = 2R': r'\frac{|AB|}{\sin(\angle APB)}=2R',
    's = ∫₀^θ √(r² + (dr/dθ)²) dθ = ∫₀^θ √((a+bθ)² + b²) dθ':
        r's=\int_0^\theta\sqrt{r^2+(dr/d\theta)^2}\,d\theta=\int_0^\theta\sqrt{(a+b\theta)^2+b^2}\,d\theta',
    's(θ) = ∫₀^θ √((a+bθ)² + b²) dθ':
        r's(\theta)=\int_0^\theta\sqrt{(a+b\theta)^2+b^2}\,d\theta',
    'min Σ(α_calculated - α_measured)²':
        r'\min\sum(\alpha_{\mathrm{calc}}-\alpha_{\mathrm{meas}})^2',
    '(xᵢ₊₁ - xᵢ)² + (yᵢ₊₁ - yᵢ)² = L²':
        r'(x_{i+1}-x_i)^2+(y_{i+1}-y_i)^2=L^2',
    'r(θ) = a + bθ': r'r(\theta)=a+b\theta',
    'r = a + bθ': r'r=a+b\theta',
    'b = h/(2π)': r'b=h/(2\pi)',
    'b = h/2π': r'b=h/2\pi',
    '∠APB = α': r'\angle APB=\alpha',
}
_FKEYS = sorted(FORMULAS, key=len, reverse=True)   # 长串优先，避免子串误匹配


def _png_size(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>II', head[16:24])


def render_math(latex):
    if latex in _mcache:
        return _mcache[latex]
    path = os.path.join(_FX, f'f{abs(hash(latex)) % 10**9}.png')
    fig = plt.figure()
    fig.text(0.5, 0.5, f'${latex}$', fontsize=13, ha='center', va='center')
    fig.savefig(path, dpi=300, transparent=True, bbox_inches='tight', pad_inches=0.02)
    plt.close(fig)
    w, h = _png_size(path)
    _mcache[latex] = (path, w, h)
    return _mcache[latex]


def tokenize_formulas(text):
    """把文本切成 [('t', 文本) | ('m', 公式纯文本)]。"""
    out, i, n = [], 0, len(text)
    while i < n:
        best = None
        for k in _FKEYS:
            idx = text.find(k, i)
            if idx != -1 and (best is None or idx < best[0]):
                best = (idx, k)
        if best is None:
            out.append(('t', text[i:]))
            break
        start, k = best
        if start > i:
            out.append(('t', text[i:start]))
        out.append(('m', k))
        i = start + len(k)
    return out


MEMBERS = [
    ('队员1_学习心得模板_建模员.md', '谢俊邦', '建模员', '负责模型构建、公式推导、假设合理性分析'),
    ('队员2_学习心得模板_编程员.md', '周俊皓', '编程员', '负责算法实现、数值计算、数据处理与可视化'),
    ('队员3_学习心得模板_写作员.md', '张霖熙', '写作员', '负责论文结构、摘要撰写、图表与排版'),
]
DATES = ['7 月 10 日', '7 月 15 日', '7 月 17 日']   # 研读日期
HOURS = ['5', '6', '2.5']                            # 用时


# ---------------- 第一步：清理 markdown ----------------
def unwrap(line):
    """剥掉行首所有 '>' 引用层，返回 (层数, 内容)。"""
    depth, s = 0, line
    while True:
        m = re.match(r'^>\s?(.*)$', s)
        if not m:
            break
        s = m.group(1)
        depth += 1
    return depth, s


def clean(raw):
    """返回 [(kind, text, quote)]，kind ∈ h1/h2/h3/table/hr/code/li/p/blank。"""
    lines = raw.split('\n')
    items = []
    i, n = 0, len(lines)
    while i < n:
        raw_line = lines[i]

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', raw_line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            text = re.sub(r'（≥\d+字）|（可选）', '', text).strip()  # 去字数/可选标记
            text = text.replace('*', '')
            items.append((f'h{min(level,3)}', text, False))
            i += 1
            continue

        # 表格行
        if raw_line.strip().startswith('|'):
            items.append(('table', raw_line, False))
            i += 1
            continue

        # 分隔线
        if raw_line.strip() == '---':
            items.append(('hr', '', False))
            i += 1
            continue

        depth, content = unwrap(raw_line)

        # 代码块
        if content.strip().startswith('```'):
            i += 1
            while i < n:
                d2, c2 = unwrap(lines[i])
                if c2.strip().startswith('```'):
                    i += 1
                    break
                items.append(('code', c2, False))
                i += 1
            continue

        if depth == 0:
            # 去掉原模板页脚（撰写日期/总字数），落款由本脚本另加
            if re.search(r'撰写日期|总字数', content):
                i += 1
                continue
            items.append(('p' if content.strip() else 'blank', content, False))
            i += 1
            continue

        # 引用块内容
        c = content.strip()
        # 斜体提示语（整行 *…*，非 **粗体**）→ 跳过该提示块直到空行
        if depth == 1 and c.startswith('*') and c.endswith('*') and not c.startswith('**'):
            i += 1
            while i < n:
                _, c2 = unwrap(lines[i])
                i += 1
                if not c2.strip():
                    break
            items.append(('blank', '', False))
            continue
        # 删掉“本文结构/主攻方向/备选方向/角色定位”这类顶部脚手架说明（已并入标题）
        if any(kw in c for kw in ('本文结构', '心得与效果分开写', '主攻方向', '备选方向', '角色定位')):
            i += 1
            continue

        quote = depth >= 2
        if not c:
            items.append(('blank', '', False))
        elif re.match(r'^(\d+\.|[-*])\s+', c):
            items.append(('li', c, quote))
        else:
            items.append(('p', c, quote))
        i += 1
    return items


# ---------------- 第二步：写 Word ----------------
def set_run(run, font='宋体', size=12, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)


EMOJI = re.compile(
    '[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F0FF'
    '\U00002190-\U000021FF\U00002B00-\U00002BFF️✅❌✔]')

_fx_count = [0]


def emit_text(p, text, size=12, italic=False, font='宋体'):
    """去掉 ** 粗体标记与 emoji，识别关键公式内嵌为排版图（降低 AI 痕迹、公式更美）。"""
    from docx.shared import Pt as _Pt
    text = text.replace('**', '').replace('`', '')
    text = EMOJI.sub('', text).strip()
    if not text:
        return
    for kind, val in tokenize_formulas(text):
        if kind == 't':
            if val:
                run = p.add_run(val)
                set_run(run, font=font, size=size, italic=italic)
        else:
            latex = FORMULAS[val]
            path, w, h = render_math(latex)
            pt = h * 72.0 / 300.0          # 图片自然高(pt)，对应 fontsize=13
            run = p.add_run()
            run.add_picture(path, height=_Pt(pt))
            _fx_count[0] += 1


def new_doc():
    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.7)
    return doc


def render_into(doc, md_path, name, role, duty):
    """把一位队员的心得内容渲染进已有的 doc（供单份/合并复用）。"""
    raw = open(md_path, encoding='utf-8').read()
    items = clean(raw)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(t.add_run('数模国赛第一阶段学习心得'), '黑体', 16, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(sub.add_run(f'{name}（{role}）'), '黑体', 12)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(info.add_run(f'主攻 B 题，备选 C 题　|　{duty}'), '楷体', 10.5)
    doc.add_paragraph()

    table_rowbuf = []

    def flush_table():
        if not table_rowbuf:
            return
        rows = [r for r in table_rowbuf if not re.match(r'^\s*\|[\s:|-]+\|\s*$', r)]
        cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        if not cells:
            table_rowbuf.clear()
            return
        is_paperlist = any('研读日期' in c for c in cells[0])
        tb = doc.add_table(rows=len(cells), cols=len(cells[0]))
        tb.style = 'Light Grid Accent 1'
        drow = 0
        for ri, row in enumerate(cells):
            for ci, val in enumerate(row):
                val = val.replace('*', '')
                if is_paperlist and ri > 0:
                    if '研读日期' in cells[0][ci] and not val:
                        val = DATES[drow] if drow < len(DATES) else ''
                    if '用时' in cells[0][ci] and not val:
                        val = HOURS[drow] if drow < len(HOURS) else ''
                cell = tb.cell(ri, ci)
                cell.paragraphs[0].text = ''
                r = cell.paragraphs[0].add_run(val)
                set_run(r, '宋体', 10.5, bold=(ri == 0))
            if is_paperlist and ri > 0:
                drow += 1
        doc.add_paragraph()
        table_rowbuf.clear()

    for kind, text, quote in items:
        if kind == 'table':
            table_rowbuf.append(text)
            continue
        else:
            flush_table()
        if kind == 'blank' or kind == 'hr' or kind == 'h1':
            continue
        if kind == 'h2':
            p = doc.add_paragraph()
            set_run(p.add_run(text), '黑体', 14, bold=True)
        elif kind == 'h3':
            p = doc.add_paragraph()
            set_run(p.add_run(text), '黑体', 12, bold=True)
        elif kind == 'code':
            p = doc.add_paragraph()
            set_run(p.add_run(text), 'Consolas', 9)
        elif kind == 'li':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            m = re.match(r'^(\d+\.)\s+(.*)$', text)
            if m:
                emit_text(p, f'{m.group(1)} {m.group(2)}', size=12, italic=quote)
            else:
                body = re.sub(r'^[-*]\s+', '', text)
                emit_text(p, f'· {body}', size=12, italic=quote)
        else:  # p
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.82)
            p.paragraph_format.space_after = Pt(3)
            emit_text(p, text, size=12, italic=quote)
    flush_table()

    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(end.add_run(f'{name}　2026 年 7 月'), '宋体', 12)
    return len(items)


def build(md_path, name, role, duty, out_path):
    _fx_count[0] = 0
    doc = new_doc()
    k = render_into(doc, md_path, name, role, duty)
    doc.save(out_path)
    return k, _fx_count[0]


def build_merged(out_path):
    """三人合订：封面 + 每人一节（分页）。"""
    _fx_count[0] = 0
    doc = new_doc()
    for _ in range(6):
        doc.add_paragraph()
    c1 = doc.add_paragraph(); c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c1.add_run('2026 高教社杯全国大学生数学建模竞赛'), '黑体', 13)
    c2 = doc.add_paragraph(); c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c2.add_run('第一阶段 · 学习心得'), '黑体', 22, bold=True)
    doc.add_paragraph()
    c3 = doc.add_paragraph(); c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c3.add_run('主攻 B 题　　备选 C 题'), '楷体', 12)
    for _ in range(3):
        doc.add_paragraph()
    for _fn, nm, rl, _d in MEMBERS:
        m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(m.add_run(f'{nm}　·　{rl}'), '宋体', 14)
    for _ in range(4):
        doc.add_paragraph()
    dd = doc.add_paragraph(); dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(dd.add_run('2026 年 7 月'), '宋体', 12)
    for fn, name, role, duty in MEMBERS:
        doc.add_page_break()
        render_into(doc, os.path.join(BASE, fn), name, role, duty)
    doc.save(out_path)
    return _fx_count[0]


if __name__ == '__main__':
    import shutil
    for fn, name, role, duty in MEMBERS:
        out = os.path.join(OUT, f'学习心得_{name}_{role}.docx')
        k, nf = build(os.path.join(BASE, fn), name, role, duty, out)
        print(f'单份 {name}_{role}  （{k} 段，公式图 {nf} 处）')
    names = '-'.join(m[1] for m in MEMBERS)
    merged = os.path.join(OUT, f'学习心得_合并_{names}.docx')
    build_merged(merged)
    print(f'合并 {os.path.basename(merged)}')
    shutil.rmtree(_FX, ignore_errors=True)
    print('完成，输出目录：', OUT)
