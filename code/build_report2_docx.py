# -*- coding: utf-8 -*-
"""学习报告2 草稿 .md → Word(.docx)。
结构：封面 + 第一部分（三份总结性报告，每人一页起）+ 第二部分（研读清单）+ 第三部分（吴孟达讲座心得）。
复用 build_reflections_docx.py 的清理/渲染思路（宋体正文/黑体标题/楷体信息行/Light Grid 表/mathtext 公式图）。
运行：py build_report2_docx.py
"""
import re
import os
import struct
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams['mathtext.fontset'] = 'cm'

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 数模/
OUT = os.path.join(BASE, '学习报告2')
os.makedirs(OUT, exist_ok=True)
MD = os.path.join(OUT, '学习报告2_草稿.md')

# ---------------- 数学公式渲染（matplotlib mathtext → 内嵌图）----------------
_FX = os.path.join(OUT, '_fx')
os.makedirs(_FX, exist_ok=True)
_mcache = {}

# 纯文本公式 → LaTeX（按草稿中精确字符串匹配，替换为排版图；未命中则保持原文）
FORMULAS = {
    'Re = ρuD_h/μ ≈ 181': r'\mathrm{Re}=\frac{\rho u D_h}{\mu}\approx 181',
    'R* = R_cond + 1/(h·A_wet·η_fin) + 1/(2ṁc_p)':
        r'R^{*}=R_{\mathrm{cond}}+\frac{1}{h\,A_{\mathrm{wet}}\,\eta_{\mathrm{fin}}}+\frac{1}{2\dot{m}c_p}',
    'Δp* = f·(L/D_h)·ρu²/2 + N·K_p·(ρu²/2)·g(r_w) + Δp_man(r_d)':
        r'\Delta p^{*}=f\frac{L}{D_h}\frac{\rho u^2}{2}+N K_p\frac{\rho u^2}{2}g(r_w)+\Delta p_{\mathrm{man}}(r_d)',
}
_FKEYS = sorted(FORMULAS, key=len, reverse=True)  # 长串优先


def _png_size(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>II', head[16:24])


def render_math(latex):
    if latex in _mcache:
        return _mcache[latex]
    path = os.path.join(_FX, f'f{abs(hash(latex)) % 10**9}.png')
    fig = plt.figure()
    fig.text(0.5, 0.5, f'${latex}$', fontsize=13, ha='center', va='center')
    fig.savefig(path, dpi=300, transparent=True, bbox_inches='tight', pad_inches=0.02)
    plt.close(fig)
    w, h = _png_size(path)
    _mcache[latex] = (path, w, h)
    return _mcache[latex]


def tokenize_formulas(text):
    out, i, n = [], 0, len(text)
    while i < n:
        best = None
        for k in _FKEYS:
            idx = text.find(k, i)
            if idx != -1 and (best is None or idx < best[0]):
                best = (idx, k)
        if best is None:
            out.append(('t', text[i:]))
            break
        start, k = best
        if start > i:
            out.append(('t', text[i:start]))
        out.append(('m', k))
        i = start + len(k)
    return out


# ---------------- 清理 markdown ----------------
def clean(raw):
    """返回 [(kind, text)]，kind ∈ h1~h5/table/code/li/p/blank/fall。"""
    lines = raw.split('\n')
    items = []
    i, n = 0, len(lines)
    while i < n:
        raw_line = lines[i]

        m = re.match(r'^(#{1,5})\s+(.*)$', raw_line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip().replace('*', '')
            items.append((f'h{level}', text))
            i += 1
            continue

        if raw_line.strip().startswith('|'):
            items.append(('table', raw_line))
            i += 1
            continue

        if raw_line.strip() == '---':
            items.append(('blank', ''))
            i += 1
            continue

        if raw_line.strip().startswith('```'):
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                items.append(('code', lines[i]))
                i += 1
            i += 1
            continue

        c = raw_line.strip()
        if not c:
            items.append(('blank', ''))
        elif re.match(r'^(\d+\.|[-*])\s+', c):
            items.append(('li', c))
        elif c.startswith('落款：'):
            items.append(('fall', c[3:]))
        else:
            items.append(('p', c))
        i += 1
    return items


# ---------------- 写 Word ----------------
EMOJI = re.compile(
    '[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F0FF'
    '\U00002190-\U000021FF\U00002B00-\U00002BFF️✅❌✔]')


def set_run(run, font='宋体', size=12, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)


def emit_text(p, text, size=12, italic=False, font='宋体'):
    text = text.replace('**', '').replace('`', '')
    text = EMOJI.sub('', text).strip()
    if not text:
        return
    for kind, val in tokenize_formulas(text):
        if kind == 't':
            if val:
                run = p.add_run(val)
                set_run(run, font=font, size=size, italic=italic)
        else:
            latex = FORMULAS[val]
            path, w, h = render_math(latex)
            pt = h * 72.0 / 300.0
            run = p.add_run()
            run.add_picture(path, height=Pt(pt))


def new_doc():
    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.7)
    return doc


def cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    c1 = doc.add_paragraph(); c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c1.add_run('2026 高教社杯全国大学生数学建模竞赛'), '黑体', 13)
    c2 = doc.add_paragraph(); c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c2.add_run('学习报告2'), '黑体', 22, bold=True)
    doc.add_paragraph()
    c3 = doc.add_paragraph(); c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c3.add_run('主攻 B 题　　备选 C 题'), '楷体', 12)
    c4 = doc.add_paragraph(); c4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(c4.add_run('队伍号：095'), '楷体', 12)
    for _ in range(3):
        doc.add_paragraph()
    for nm, rl in [('谢俊邦', '建模员'), ('周俊皓', '编程员'), ('张霖熙', '写作员')]:
        m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(m.add_run(f'{nm}　·　{rl}'), '宋体', 14)
    for _ in range(4):
        doc.add_paragraph()
    dd = doc.add_paragraph(); dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(dd.add_run('2026 年 8 月'), '宋体', 12)


def render_md(doc, md_path):
    raw = open(md_path, encoding='utf-8').read()
    items = clean(raw)

    table_rowbuf = []

    def flush_table():
        if not table_rowbuf:
            return
        rows = [r for r in table_rowbuf if not re.match(r'^\s*\|[\s:|-]+\|\s*$', r)]
        cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        if not cells:
            table_rowbuf.clear()
            return
        tb = doc.add_table(rows=len(cells), cols=len(cells[0]))
        tb.style = 'Light Grid Accent 1'
        for ri, row in enumerate(cells):
            for ci, val in enumerate(row):
                val = val.replace('*', '')
                cell = tb.cell(ri, ci)
                cell.paragraphs[0].text = ''
                r = cell.paragraphs[0].add_run(val)
                set_run(r, '宋体', 10.5, bold=(ri == 0))
        doc.add_paragraph()
        table_rowbuf.clear()

    first_p_after_h3 = False
    prev_kind = None
    for kind, text in items:
        if kind == 'table':
            table_rowbuf.append(text)
            continue
        else:
            flush_table()
        if kind == 'blank' or kind == 'h1':
            continue
        if kind == 'h2':
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(12)
            set_run(p.add_run(text), '黑体', 15, bold=True)
            prev_kind = kind
            continue
        if kind == 'h3':
            p = doc.add_paragraph()
            if prev_kind != 'h2':  # 紧跟部分标题则同页，否则另起一页
                p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            set_run(p.add_run(text), '黑体', 14, bold=True)
            first_p_after_h3 = True
            prev_kind = kind
            continue
        if kind == 'h4':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            set_run(p.add_run(text), '黑体', 12, bold=True)
            continue
        if kind == 'h5':
            p = doc.add_paragraph()
            set_run(p.add_run(text), '宋体', 12, bold=True)
            continue
        if kind == 'code':
            p = doc.add_paragraph()
            set_run(p.add_run(text), 'Consolas', 9)
            continue
        if kind == 'fall':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_run(p.add_run(text), '宋体', 12)
            continue
        if kind == 'li':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            m = re.match(r'^(\d+\.)\s+(.*)$', text)
            if m:
                emit_text(p, f'{m.group(1)} {m.group(2)}', size=12)
            else:
                emit_text(p, f'· {re.sub(r"^[-*]\s+", "", text)}', size=12)
            continue
        # p
        p = doc.add_paragraph()
        if first_p_after_h3:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit_text(p, text, size=10.5, font='楷体')
            first_p_after_h3 = False
        else:
            p.paragraph_format.first_line_indent = Cm(0.82)
            p.paragraph_format.space_after = Pt(3)
            emit_text(p, text, size=12)
    flush_table()


if __name__ == '__main__':
    import shutil
    doc = new_doc()
    cover(doc)
    render_md(doc, MD)
    out = os.path.join(OUT, '095学习报告2.docx')
    doc.save(out)
    shutil.rmtree(_FX, ignore_errors=True)
    print('saved:', out)
